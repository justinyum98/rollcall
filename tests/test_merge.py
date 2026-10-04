import threading
from datetime import datetime

import docx
import pytest

from conftest import docx_text
from rollcall import merge
from rollcall.merge import (
    default_filename_pattern,
    generate,
    inspect_template,
    make_filename,
    make_output_folder,
)
from rollcall.roster import Column, Roster


def roster_of(*rows, columns=("First Name", "Last Name")):
    cols = [Column(label=c, key=c.lower().replace(" ", "_")) for c in columns]
    return Roster(columns=cols, rows=[dict(zip([c.key for c in cols], r)) for r in rows])


# --- file names -------------------------------------------------------------


def test_default_filename_pattern():
    assert default_filename_pattern(roster_of()) == "{last_name}, {first_name}"
    other = roster_of(columns=("Student", "Grade"))
    assert default_filename_pattern(other) == "{student}"


@pytest.mark.parametrize(
    "pattern, row, expected",
    [
        ("{last_name}, {first_name}", {"first_name": "Ada", "last_name": "Lovelace"}, "Lovelace, Ada"),
        ("{first_name} - {grade}", {"first_name": "Ada"}, "Ada -"),
        ("{name}", {"name": 'A/B\\C:D*E?F"G<H>I|J'}, "ABCDEFGHIJ"),
        ("{name}", {"name": "trailing dots... "}, "trailing dots"),
        ("{name}", {"name": "CON"}, "CON_"),
        ("{name}", {"name": "nul"}, "nul_"),
        ("{name}", {"name": ""}, "Student"),
        ("{name}", {"name": "José 李"}, "José 李"),
        ("Certificate", {}, "Certificate"),
    ],
)
def test_make_filename(pattern, row, expected):
    assert make_filename(pattern, row) == expected


def test_make_filename_is_capped():
    assert len(make_filename("{n}", {"n": "x" * 500})) == 120


def test_make_output_folder_never_reuses(tmp_path):
    now = datetime(2026, 10, 3, 15, 30)
    first = make_output_folder(tmp_path, tmp_path / "Report: Card.docx", now)
    second = make_output_folder(tmp_path, tmp_path / "Report: Card.docx", now)
    assert first.name == "Report Card - 2026-10-03 1530"
    assert second.name == "Report Card - 2026-10-03 1530 (2)"
    assert first.is_dir() and second.is_dir()


# --- template checks --------------------------------------------------------


def test_inspect_all_match(make_docx):
    path = make_docx(["Dear {{first_name}} {{ last_name }},"])
    check = inspect_template(path, roster_of())
    assert check.ok
    assert check.placeholders == {"first_name", "last_name"}
    assert check.unknown == []
    assert check.error is None


def test_inspect_accepts_column_labels(make_docx):
    path = make_docx(["Dear {{First Name}} {{LAST NAME}},"])
    check = inspect_template(path, roster_of())
    assert check.ok
    assert check.placeholders == {"first_name", "last_name"}


def test_inspect_reports_unknown_and_unused(make_docx):
    path = make_docx(["Dear {{first_name}}, you are in {{grade}}."])
    check = inspect_template(path, roster_of())
    assert not check.ok
    assert check.unknown == ["grade"]
    assert check.unused == ["last_name"]


def test_inspect_unknown_label_with_spaces(make_docx):
    path = make_docx(["Dear {{First Name}}, homeroom {{Home Room}}."])
    check = inspect_template(path, roster_of())
    assert not check.ok
    assert check.unknown == ["Home Room"]
    assert check.error is None


def test_inspect_broken_placeholder(make_docx):
    path = make_docx(["Dear {{first_name}, welcome {% if %}"])
    check = inspect_template(path, roster_of())
    assert not check.ok
    assert "{{first_name}" in check.error


@pytest.mark.parametrize(
    "text, fragment",
    [
        ("Dear {first_name}}, welcome", "{first_name}}"),
        ("Dear {{ first_name }} {{{last_name}}}", "{{{last_name}}}"),
    ],
)
def test_inspect_broken_but_valid_jinja(make_docx, text, fragment):
    check = inspect_template(make_docx([text]), roster_of())
    assert not check.ok
    assert fragment in check.error


def test_inspect_ignores_ordinary_braces(make_docx):
    check = inspect_template(make_docx(["{{first_name}} picked {1, 2, 3}", "{{last_name}}"]), roster_of())
    assert check.ok


def test_inspect_finds_no_placeholders(make_docx):
    check = inspect_template(make_docx(["Hello everyone"]), roster_of())
    assert not check.ok
    assert "no placeholders" in check.error


def test_inspect_header_and_footer(tmp_path):
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "{{ grade }}"
    document.add_paragraph("{{ first_name }}")
    path = tmp_path / "t.docx"
    document.save(path)
    assert inspect_template(path, roster_of()).unknown == ["grade"]


def test_inspect_not_a_docx(tmp_path):
    path = tmp_path / "t.docx"
    path.write_text("not a zip")
    check = inspect_template(path, roster_of())
    assert not check.ok and "Word" in check.error


# --- generating documents ---------------------------------------------------


def test_generate_creates_one_doc_per_student(tmp_path, make_docx):
    template = make_docx(["Certificate for ", ["{{first", "_name}}"], "Family: {{ Last Name }}"])
    roster = roster_of(("Ada", "Lovelace"), ("José", "Núñez"), ("Ada", "Lovelace"))
    progress = []

    result = generate(
        roster,
        template,
        tmp_path / "out",
        "{first_name} {last_name}",
        on_progress=lambda done, total, name: progress.append((done, total, name)),
    )

    assert result.errors == []
    assert not result.cancelled
    assert [p.name for p in result.created] == [
        "Ada Lovelace.docx",
        "José Núñez.docx",
        "Ada Lovelace (2).docx",
    ]
    assert all(p.parent == result.folder for p in result.created)
    assert docx_text(result.created[1]) == "Certificate for \nJosé\nFamily: Núñez"
    assert progress[-1] == (3, 3, "Ada Lovelace")


def test_generate_keeps_formatting(tmp_path, make_docx):
    template = make_docx([["Hello ", "{{first_name}}", "!"]])
    result = generate(roster_of(("Ada", "L")), template, tmp_path, "{first_name}")
    runs = docx.Document(result.created[0]).paragraphs[0].runs
    bold = [r.text for r in runs if r.bold]
    assert bold == ["Ada"]


def test_generate_continues_after_a_failure(tmp_path, make_docx, monkeypatch):
    template = make_docx(["{{first_name}}"])
    real_render = merge._render_one

    def flaky(template_bytes, keys, row, dest):
        if row["first_name"] == "Bad":
            raise RuntimeError("disk full")
        return real_render(template_bytes, keys, row, dest)

    monkeypatch.setattr(merge, "_render_one", flaky)
    roster = roster_of(("Ada", "L"), ("Bad", "Row"), ("Alan", "T"))
    result = generate(roster, template, tmp_path, "{first_name}")
    assert [p.stem for p in result.created] == ["Ada", "Alan"]
    assert result.errors == [("Bad Row", "disk full")]


def test_generate_can_be_cancelled(tmp_path, make_docx):
    template = make_docx(["{{first_name}}"])
    cancel = threading.Event()
    roster = roster_of(("A", "1"), ("B", "2"), ("C", "3"))

    def stop_after_first(done, total, name):
        cancel.set()

    result = generate(roster, template, tmp_path, "{first_name}", stop_after_first, cancel)
    assert result.cancelled
    assert len(result.created) == 1


def test_generate_escapes_special_characters(tmp_path, make_docx):
    template = make_docx(["Family: {{last_name}}"])
    roster = roster_of(("Ada", "Smith & <Jones>"))
    result = generate(roster, template, tmp_path, "{last_name}")
    assert result.errors == []
    assert docx_text(result.created[0]) == "Family: Smith & <Jones>"
    assert result.created[0].name == "Smith & Jones.docx"

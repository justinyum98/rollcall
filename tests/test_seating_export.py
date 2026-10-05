import docx
from docx.enum.section import WD_ORIENT

from rollcall.seating.export import save_chart_docx
from rollcall.seating.model import Layout, SeatingClass, Student


def make_class(layout, n):
    students = [
        Student(f"First{i}", f"Last{i}", behavior="loud", needs_front=True, special_needs=True, tolerant=True)
        for i in range(n)
    ]
    cls = SeatingClass(name="Period 3", students=students, layout=layout)
    cls.chart = {seat.id: s.id for seat, s in zip(layout.seats(), students)}
    return cls


def grid(path):
    table = docx.Document(path).tables[0]
    return [[cell.text for cell in row.cells] for row in table.rows]


def all_text(path):
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    parts += [c.text for t in d.tables for row in t.rows for c in row.cells]
    return "\n".join(parts)


def test_rows_layout(tmp_path):
    cls = make_class(Layout("rows", rows=2, cols=3), 5)
    path = save_chart_docx(cls, tmp_path / "chart.docx")
    assert grid(path) == [
        ["First0 Last0", "First1 Last1", "First2 Last2"],
        ["First3 Last3", "First4 Last4", ""],
    ]
    d = docx.Document(path)
    assert d.sections[0].orientation == WD_ORIENT.LANDSCAPE
    text = all_text(path)
    assert "Period 3" in text and "Front of room" in text


def test_pairs_layout_has_aisles(tmp_path):
    cls = make_class(Layout("pairs", rows=1, cols=2), 4)
    path = save_chart_docx(cls, tmp_path / "chart.docx")
    assert grid(path) == [["First0 Last0", "First1 Last1", "", "First2 Last2", "First3 Last3"]]


def test_tables_layout_lists_each_table(tmp_path):
    cls = make_class(Layout("tables", rows=1, cols=2, table_size=3), 5)
    path = save_chart_docx(cls, tmp_path / "chart.docx")
    cells = grid(path)[0]
    assert cells[0] == "Table 1\nFirst0 Last0\nFirst1 Last1\nFirst2 Last2"
    assert cells[1] == ""  # space between tables
    assert cells[2] == "Table 2\nFirst3 Last3\nFirst4 Last4"


def test_printout_never_shows_classifications(tmp_path):
    for layout in (Layout("rows", rows=2, cols=3), Layout("pairs", rows=2, cols=2), Layout("tables", rows=1, cols=2, table_size=3)):
        text = all_text(save_chart_docx(make_class(layout, 5), tmp_path / f"{layout.kind}.docx")).lower()
        for word in ("loud", "quiet", "special", "tolerant", "vision", "needs"):
            assert word not in text, (layout.kind, word)


def test_unseated_students_are_listed(tmp_path):
    cls = make_class(Layout("rows", rows=1, cols=2), 2)
    cls.students.append(Student("New", "Kid"))
    text = all_text(save_chart_docx(cls, tmp_path / "chart.docx"))
    assert "Not seated yet: New Kid" in text

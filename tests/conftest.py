import csv
from pathlib import Path

import docx
import openpyxl
import pytest


@pytest.fixture
def make_xlsx(tmp_path):
    def _make(rows, name="students.xlsx"):
        wb = openpyxl.Workbook()
        ws = wb.active
        for row in rows:
            ws.append(row)
        path = tmp_path / name
        wb.save(path)
        return path

    return _make


@pytest.fixture
def make_csv(tmp_path):
    def _make(rows, name="students.csv", encoding="utf-8-sig"):
        path = tmp_path / name
        with open(path, "w", newline="", encoding=encoding) as f:
            csv.writer(f).writerows(rows)
        return path

    return _make


@pytest.fixture
def make_docx(tmp_path):
    """Build a .docx where each paragraph is a list of runs.

    A paragraph given as a plain string becomes a single run. A list of
    strings becomes one run per string, which mimics how Word splits a
    placeholder like {{first_name}} across formatting runs.
    """

    def _make(paragraphs, name="template.docx"):
        document = docx.Document()
        for para in paragraphs:
            runs = [para] if isinstance(para, str) else para
            p = document.add_paragraph()
            for i, text in enumerate(runs):
                run = p.add_run(text)
                run.bold = i % 2 == 1
        path = tmp_path / name
        document.save(path)
        return path

    return _make


def docx_text(path: Path) -> str:
    return "\n".join(p.text for p in docx.Document(path).paragraphs)

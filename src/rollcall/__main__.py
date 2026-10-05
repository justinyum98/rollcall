"""Start RollCall. Absolute imports, because PyInstaller runs this file as a script."""

import os
import sys


def _ensure_std_streams() -> None:
    # A windowed app (no console) has sys.stdout/sys.stderr set to None, and
    # any library that prints would crash. Send that output nowhere instead.
    for name in ("stdout", "stderr"):
        if getattr(sys, name) is None:
            setattr(sys, name, open(os.devnull, "w"))


def self_test() -> int:
    """Check that a built app works end to end without anyone clicking. Used by CI."""
    import tempfile
    from pathlib import Path

    import docx
    import openpyxl

    from rollcall.gui import RollCallApp
    from rollcall.merge import generate, inspect_template
    from rollcall.roster import load_roster
    from rollcall.seating.export import save_chart_docx
    from rollcall.seating.model import Layout, SeatingClass, Student
    from rollcall.seating.solver import arrange
    from rollcall.starter import create_starter_template

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        wb = openpyxl.Workbook()
        wb.active.append(["First Name", "Last Name"])
        wb.active.append(["José", "Núñez"])
        wb.save(tmp / "students.xlsx")

        roster = load_roster(tmp / "students.xlsx")
        template = create_starter_template(roster, tmp / "template.docx")
        check = inspect_template(template, roster)
        assert check.ok, check
        result = generate(roster, template, tmp / "out", "{last_name}, {first_name}")
        assert result.errors == [], result.errors
        text = "\n".join(p.text for p in docx.Document(result.created[0]).paragraphs)
        assert "First Name: José" in text, text

        # Seating: arrange a class with a rule, then export the printable chart.
        cls = SeatingClass(name="Self-test", layout=Layout("pairs", rows=2, cols=2))
        cls.students = [Student("Ada", "L", behavior="loud"), Student("Alan", "T", behavior="loud"),
                        Student("Grace", "H", behavior="quiet"), Student("Mei", "C")]
        cls.keep_apart = [(cls.students[0].id, cls.students[1].id)]
        arrangement = arrange(cls, seed=1)
        assert arrangement.problems == [], arrangement.problems
        cls.chart = arrangement.chart
        chart = save_chart_docx(cls, tmp / "chart.docx")
        assert "Ada L" in "".join(c.text for t in docx.Document(chart).tables for r in t.rows for c in r.cells)

    # Build the window once, so missing theme or font files fail here too.
    app = RollCallApp()
    app.update()
    app.destroy()
    return 0


def main() -> None:
    _ensure_std_streams()
    if "--self-test" in sys.argv:
        try:
            sys.exit(self_test())
        except Exception as e:
            print(f"Self-test failed: {e!r}", file=sys.stderr)
            sys.exit(1)

    from rollcall.gui import run

    run()


if __name__ == "__main__":
    main()

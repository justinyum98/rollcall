"""Build a ready-to-edit template from a spreadsheet's columns."""

from pathlib import Path

import docx
from docx.shared import Pt, RGBColor

from .roster import Column, Roster, make_key


def placeholder_for(column: Column) -> str:
    """The placeholder a teacher should type for a column.

    Prefer the header as written ({{First Name}}) since that's what the
    teacher recognizes. A repeated header would point at the first column
    with that name, so those fall back to the unique key ({{name_2}}).
    """
    if make_key(column.label) == column.key:
        return "{{%s}}" % column.label
    return "{{%s}}" % column.key


def create_starter_template(roster: Roster, dest: str | Path) -> Path:
    dest = Path(dest)
    document = docx.Document()
    document.add_heading("My template", level=1)

    note = document.add_paragraph().add_run(
        "Edit this document however you like: change the wording, fonts, and layout, "
        "add a logo, and so on. Each name in double curly braces below is replaced "
        "with that student's information. You can move, copy, or delete them. "
        "Delete this note when you're done."
    )
    note.italic = True
    note.font.size = Pt(10)
    note.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    for column in roster.columns:
        document.add_paragraph(f"{column.label}: {placeholder_for(column)}")

    document.save(dest)
    return dest

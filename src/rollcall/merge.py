"""Fill a Word template once per student."""

import html
import re
import threading
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Callable

import jinja2
from docx.oxml.ns import qn
from docxtpl import DocxTemplate

from .roster import Roster, make_key

ProgressCallback = Callable[[int, int, str], None]

# A complete {{ ... }} tag inside one piece of text. Excluding "<" keeps a
# broken tag from matching across Word's XML into a later paragraph.
_TAG = re.compile(r"\{\{([^{}<]*)\}\}")
# Anything brace-shaped, used to point at a broken placeholder.
_BRACES = re.compile(r"\{+[^{}\n]*\}*")
_GOOD_TAG = re.compile(r"\{\{[^{}]*\}\}")
_ILLEGAL_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_WINDOWS_RESERVED = {"con", "prn", "aux", "nul"} | {f"{p}{i}" for p in ("com", "lpt") for i in range(1, 10)}
_MAX_FILENAME = 120


class _Template(DocxTemplate):
    """A DocxTemplate that also accepts column labels as placeholders.

    Teachers naturally type {{First Name}} to match their spreadsheet, but
    Jinja only understands {{ first_name }}. docxtpl's patch_xml() runs on
    every part of the document after Word's split runs have been merged, so
    it's the one place to rewrite labels into keys.
    """

    def __init__(self, template_file, keys):
        self._keys = set(keys)
        self.unmatched_labels: set[str] = set()
        super().__init__(template_file)

    def patch_xml(self, src_xml):
        return _TAG.sub(self._rewrite_tag, super().patch_xml(src_xml))

    def _rewrite_tag(self, match: re.Match) -> str:
        text = html.unescape(match.group(1)).strip()
        key = make_key(text)
        if key in self._keys:
            return "{{ %s }}" % key
        if not _is_jinja_expression(text):
            # A label with spaces that matches no column, e.g. {{Home Room}}.
            # Render it as blank instead of failing the whole document.
            self.unmatched_labels.add(text)
            return "{{ '' }}"
        return match.group(0)


def _is_jinja_expression(text: str) -> bool:
    try:
        jinja2.Environment().parse("{{ %s }}" % text)
        return True
    except jinja2.TemplateSyntaxError:
        return False


# --- template checks --------------------------------------------------------


@dataclass
class TemplateCheck:
    placeholders: set[str] = field(default_factory=set)  # matched column keys
    unknown: list[str] = field(default_factory=list)  # in template, not in spreadsheet
    unused: list[str] = field(default_factory=list)  # in spreadsheet, not in template
    error: str | None = None  # the template can't be used at all

    @property
    def ok(self) -> bool:
        return self.error is None and not self.unknown


def inspect_template(path: str | Path, roster: Roster) -> TemplateCheck:
    path = Path(path)
    try:
        tpl = _Template(str(path), roster.keys)
        broken = _broken_placeholders(tpl.get_docx())
        variables = tpl.get_undeclared_template_variables()
    except jinja2.TemplateSyntaxError:
        return TemplateCheck(error=_broken_message(broken))
    except Exception as e:
        return TemplateCheck(
            error=f'Couldn\'t open "{path.name}" as a Word document. '
            "Make sure it's a .docx file and isn't open in another program. "
            f"(Details: {type(e).__name__}: {e})"
        )
    if broken:
        # Not a Jinja error (e.g. "{first_name}}"), but it would show up as
        # literal braces in every student's document.
        return TemplateCheck(error=_broken_message(broken))

    found = variables & set(roster.keys)
    unknown = sorted(variables - found) + sorted(tpl.unmatched_labels)
    if not found and not unknown:
        return TemplateCheck(
            error="This template has no placeholders yet. Type a placeholder such as "
            "{{First Name}} where each student's information should go."
        )
    return TemplateCheck(
        placeholders=found,
        unknown=unknown,
        unused=[k for k in roster.keys if k not in found],
    )


def _broken_placeholders(document) -> list[str]:
    """Brace fragments that look like a placeholder attempt but aren't {{like_this}}.

    Only fragments containing "{{" or "}}" count, so ordinary text such as
    "{1, 2, 3}" is left alone.
    """
    return [
        fragment
        for text in _paragraph_texts(document)
        for fragment in _BRACES.findall(text)
        if ("{{" in fragment or "}}" in fragment) and not _GOOD_TAG.fullmatch(fragment)
    ]


def _broken_message(broken: list[str]) -> str:
    advice = "Each placeholder needs two curly braces on each side, like {{First Name}}."
    if broken:
        return f"This placeholder looks broken: {broken[0].strip()}\n{advice}"
    return f"A placeholder in the template is broken. {advice}"


def _paragraph_texts(document) -> list[str]:
    """Text of every paragraph in the body (including tables), headers and footers."""
    # Skip headers/footers that don't exist: reading one would make python-docx
    # create it, which also fails inside the packaged app.
    parts = [document.element.body] + [
        hf._element
        for section in document.sections
        for hf in (section.header, section.footer)
        if not hf.is_linked_to_previous
    ]
    # Join the w:t text nodes directly: lxml's itertext() repeats each run's
    # text on python-docx's custom element classes.
    return [
        "".join(t.text or "" for t in p.iter(qn("w:t")))
        for part in parts
        for p in part.iter(qn("w:p"))
    ]


# --- file names -------------------------------------------------------------


def default_filename_pattern(roster: Roster) -> str:
    if {"first_name", "last_name"} <= set(roster.keys):
        return "{last_name}, {first_name}"
    return "{%s}" % roster.keys[0]


def make_filename(pattern: str, row: dict[str, str]) -> str:
    """Fill a pattern like "{last_name}, {first_name}" and make it safe as a file name."""
    name = re.sub(r"\{(\w+)\}", lambda m: row.get(m.group(1), ""), pattern)
    return _safe_name(name) or "Student"


def _safe_name(name: str) -> str:
    name = _ILLEGAL_FILENAME_CHARS.sub("", name)
    name = re.sub(r"\s+", " ", name).strip().rstrip(". ")
    name = name[:_MAX_FILENAME].rstrip(". ")
    if name.lower() in _WINDOWS_RESERVED:
        name += "_"
    return name


def _unique_path(folder: Path, stem: str, suffix: str) -> Path:
    path, n = folder / f"{stem}{suffix}", 2
    while path.exists():
        path, n = folder / f"{stem} ({n}){suffix}", n + 1
    return path


def make_output_folder(out_dir: str | Path, template_path: str | Path, now: datetime | None = None) -> Path:
    """Create a new folder for this run so earlier results are never overwritten."""
    now = now or datetime.now()
    stem = _safe_name(Path(template_path).stem) or "RollCall"
    base = f"{stem} - {now:%Y-%m-%d %H%M}"
    folder, n = Path(out_dir) / base, 2
    while True:
        try:
            folder.mkdir(parents=True)
            return folder
        except FileExistsError:
            folder, n = Path(out_dir) / f"{base} ({n})", n + 1


# --- generating documents ---------------------------------------------------


@dataclass
class MergeResult:
    folder: Path
    created: list[Path] = field(default_factory=list)
    errors: list[tuple[str, str]] = field(default_factory=list)  # (student, message)
    cancelled: bool = False


def generate(
    roster: Roster,
    template_path: str | Path,
    out_dir: str | Path,
    filename_pattern: str,
    on_progress: ProgressCallback | None = None,
    cancel_event: threading.Event | None = None,
) -> MergeResult:
    template_bytes = Path(template_path).read_bytes()
    result = MergeResult(folder=make_output_folder(out_dir, template_path))
    total = len(roster.rows)

    for i, row in enumerate(roster.rows, start=1):
        if cancel_event is not None and cancel_event.is_set():
            result.cancelled = True
            break
        name = roster.display_name(row)
        dest = _unique_path(result.folder, make_filename(filename_pattern, row), ".docx")
        try:
            _render_one(template_bytes, roster.keys, row, dest)
            result.created.append(dest)
        except Exception as e:
            result.errors.append((name, str(e)))
        if on_progress:
            on_progress(i, total, name)

    return result


def _render_one(template_bytes: bytes, keys: list[str], row: dict[str, str], dest: Path) -> None:
    # A DocxTemplate can only be rendered once, so build a fresh one per student.
    tpl = _Template(BytesIO(template_bytes), keys)
    # autoescape: a name like "Smith & Jones" must not break the document's XML.
    tpl.render(row, autoescape=True)
    tpl.save(dest)

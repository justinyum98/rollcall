"""Print-ready seating chart as a Word document. Names only, never classifications."""

from datetime import date
from pathlib import Path

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .model import SeatingClass

_GRAY = RGBColor(0x66, 0x66, 0x66)


def save_chart_docx(cls: SeatingClass, path: str | Path) -> Path:
    path = Path(path)
    document = docx.Document()
    _landscape(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"{cls.name} · Seating chart")
    run.bold, run.font.size = True, Pt(20)
    today = date.today()
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"{today:%B} {today.day}, {today.year}")
    run.font.size, run.font.color.rgb = Pt(10), _GRAY

    front = document.add_paragraph()
    front.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = front.add_run("Front of room")
    run.bold, run.font.size = True, Pt(12)
    _shade_paragraph(front, "D9D9D9")

    names = {s.id: s.name for s in cls.students}
    if cls.layout.kind == "tables":
        _tables(document, cls, names)
    else:
        _desks(document, cls, names)

    unseated = [s.name for s in cls.students if s.id not in cls.chart.values()]
    if unseated:
        note = document.add_paragraph()
        run = note.add_run("Not seated yet: " + ", ".join(unseated))
        run.font.size, run.font.color.rgb = Pt(10), _GRAY

    document.save(path)
    return path


def _desks(document, cls: SeatingClass, names: dict[str, str]) -> None:
    """Rows of single desks, or rows of desk pairs with an empty aisle column between pairs."""
    layout = cls.layout
    seat_cols = layout.cols * 2 if layout.kind == "pairs" else layout.cols
    # Map each seat column to a table column, skipping one column after every pair.
    to_col = {c: c + (c // 2 if layout.kind == "pairs" else 0) for c in range(seat_cols)}
    width = seat_cols + (layout.cols - 1 if layout.kind == "pairs" else 0)

    table = document.add_table(rows=layout.rows, cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    size = Pt(11 if width <= 8 else 9)
    for seat in layout.seats():
        cell = table.cell(seat.row, to_col[seat.col])
        _border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        name = names.get(cls.chart.get(seat.id), "")
        _cell_text(cell, [name], size)
    for row in table.rows:
        row.height = Inches(0.6)
    _set_widths(table, gaps=[c for c in range(width) if c not in to_col.values()], gap_width=0.35)


def _tables(document, cls: SeatingClass, names: dict[str, str]) -> None:
    """Each table as one box listing its students, with space between tables."""
    layout = cls.layout
    width, height = layout.cols * 2 - 1, layout.rows * 2 - 1
    grid = document.add_table(rows=height, cols=width)
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    by_table: dict[tuple[int, int], list[str]] = {}
    for seat in layout.seats():
        name = names.get(cls.chart.get(seat.id))
        by_table.setdefault((seat.row, seat.col), [])
        if name:
            by_table[(seat.row, seat.col)].append(name)
    for n, ((r, c), people) in enumerate(sorted(by_table.items()), start=1):
        cell = grid.cell(r * 2, c * 2)
        _border(cell)
        _cell_text(cell, [f"Table {n}", *people], Pt(11), bold_first=True)
    _set_widths(grid, gaps=list(range(1, width, 2)), gap_width=0.5)


# --- low-level Word formatting ------------------------------------------------


def _landscape(document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(0.5))


def _cell_text(cell, lines: list[str], size, bold_first: bool = False) -> None:
    for i, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # The default style adds space after each paragraph, which pushes
        # names toward the top of the desk instead of the middle.
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(line)
        run.font.size = size
        run.bold = bold_first and i == 0


def _border(cell) -> None:
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "8"), ("color", "808080")):
            element.set(qn(f"w:{key}"), value)
        borders.append(element)
    cell._tc.get_or_add_tcPr().append(borders)


_PRINTABLE_WIDTH = 10.0  # inches: landscape letter minus half-inch margins


def _set_widths(table, gaps: list[int], gap_width: float) -> None:
    """Narrow aisle/gap columns; the rest share the page width equally.

    Word spreads columns evenly unless every cell has an explicit width,
    which made aisles as wide as desks and wrapped students' names.
    """
    table.autofit = False
    count = len(table.columns)
    gaps = set(gaps)
    desk_width = (_PRINTABLE_WIDTH - gap_width * len(gaps)) / max(1, count - len(gaps))
    for c, column in enumerate(table.columns):
        width = Inches(gap_width if c in gaps else desk_width)
        column.width = width
        for cell in column.cells:
            cell.width = width


def _shade_paragraph(paragraph, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shading)

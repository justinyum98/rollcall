"""Regenerate the sample files in examples/. Run: uv run python scripts/make_examples.py"""

from datetime import date
from pathlib import Path

import docx
import openpyxl
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

EXAMPLES = Path(__file__).resolve().parent.parent / "examples"

STUDENTS = [
    ("Ada", "Lovelace", 5), ("Alan", "Turing", 5), ("Grace", "Hopper", 4), ("José", "Núñez", 4),
    ("Mei", "Chen", 5), ("Priya", "Patel", 3), ("Liam", "O'Brien", 3), ("Zoë", "Müller", 4),
    ("Kwame", "Mensah", 5), ("Sofia", "García", 3), ("Noah", "Kim", 4), ("Aiyana", "Begay", 5),
    ("Omar", "Haddad", 3), ("Hana", "Sato", 4), ("Lucas", "Silva", 5), ("Emma", "Johansson", 3),
    ("Diego", "Hernández", 4), ("Fatima", "Al-Sayed", 5), ("Ivan", "Petrov", 3), ("Chloé", "Dubois", 4),
    ("Tariq", "Williams", 5), ("Nia", "Okafor", 3), ("Ethan", "Nguyen", 4), ("Maya", "Cohen", 5),
    ("Leo", "Rossi", 3),
]


def make_roster():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Class list"
    ws.append(["First Name", "Last Name", "Grade", "Teacher", "Date"])
    for first, last, grade in STUDENTS:
        ws.append([first, last, grade, "Ms. Rivera", date(2026, 6, 12)])
    for col, width in zip("ABCDE", (14, 14, 8, 14, 14)):
        ws.column_dimensions[col].width = width
    wb.save(EXAMPLES / "Class list.xlsx")


# (behavior, vision, special needs, tolerant, not next to) for the seating example.
SEATING = {
    "Ada": ("loud", "", "", "", "Alan Turing"),
    "Alan": ("loud", "", "", "", ""),
    "Grace": ("quiet", "", "", "yes", ""),
    "José": ("", "x", "", "", ""),
    "Mei": ("quiet", "", "", "yes", ""),
    "Priya": ("", "x", "", "yes", ""),
    "Liam": ("LOUD", "", "Yes", "", ""),
    "Zoë": ("quiet", "", "", "yes", ""),
    "Kwame": ("", "", "", "yes", "Omar Haddad"),
    "Sofia": ("loud", "", "", "", ""),
    "Noah": ("", "x", "", "", ""),
    "Aiyana": ("quiet", "", "", "yes", ""),
    "Omar": ("", "", "", "", ""),
    "Hana": ("quiet", "", "", "yes", ""),
    "Lucas": ("loud", "", "", "", "Sofia García"),
    "Emma": ("", "", "yes", "", ""),
    "Diego": ("", "", "", "yes", ""),
    "Fatima": ("quiet", "", "", "", ""),
}


def make_seating_example():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Seating"
    ws.append(["First Name", "Last Name", "Behavior", "Vision", "Special Needs", "Tolerant", "Not Next To"])
    for first, last, _ in STUDENTS[:24]:
        ws.append([first, last, *SEATING.get(first, ("", "", "", "", ""))])
    for col, width in zip("ABCDEFG", (12, 14, 10, 8, 14, 10, 18)):
        ws.column_dimensions[col].width = width
    wb.save(EXAMPLES / "Seating example.xlsx")


def make_certificate():
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "Room 12 · {{Teacher}}"
    blue, orange = RGBColor(0x1F, 0x4E, 0x79), RGBColor(0xC0, 0x50, 0x00)

    add_line(document, [("Certificate of Achievement", 32, True, blue)], space_after=24)
    add_line(document, [("This certificate is proudly presented to", 14, False, None)])
    add_line(document, [("{{First Name}} {{Last Name}}", 28, True, orange)], space_after=18)
    add_line(document, [
        ("for outstanding effort and a wonderful year in grade ", 14, False, None),
        ("{{Grade}}", 14, True, None),
        (".", 14, False, None),
    ])
    add_line(document, [("Keep shining, {{First Name}}!", 14, False, None)], space_after=48)
    add_line(document, [("{{Teacher}}  ·  {{Date}}", 12, False, None)])
    document.save(EXAMPLES / "Certificate template.docx")


def add_line(document, runs, space_after=12):
    """A centered paragraph made of (text, size, bold, color) runs."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(space_after)
    for text, size, bold, color in runs:
        run = paragraph.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        if color is not None:
            run.font.color.rgb = color


if __name__ == "__main__":
    EXAMPLES.mkdir(exist_ok=True)
    make_roster()
    make_certificate()
    make_seating_example()
    print(f"Wrote examples to {EXAMPLES}")
